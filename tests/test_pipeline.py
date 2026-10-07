"""Business invariants and end-to-end buyer artifact checks."""

from __future__ import annotations

from datetime import date
from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from src.config import Scenario
from src.data_generator import generate
from src.data_io import normalize_and_validate
from src.po_exporter import excel_bytes, result_tables
from src.replenishment_engine import (
    _open_purchase_supply,
    decide,
    prepare_forecaster,
    round_purchase,
    valid_purchase,
)
from src.simulation import replay
from src.sql_staging import reconcile_ledger, stage_position


@pytest.fixture(scope="module")
def sample():
    generated = generate(42)
    data, issues = normalize_and_validate({k: v for k, v in generated.items() if k != "truth_demand"})
    assert issues.empty
    return generated, data, prepare_forecaster(data)


def test_ledger_and_position(sample):
    _, data, _ = sample
    assert reconcile_ledger(data).empty
    position = stage_position(data, date(2026, 9, 18))
    expected = (
        position["on_hand_units"]
        + position["supplier_on_order_units"]
        + position["transfer_in_transit_units"]
        - position["allocated_units"]
    )
    assert (position["inventory_position"] == expected).all()
    assert len(position) == 24 * 4


def test_bad_snapshot_and_missing_sale_are_detected(sample):
    _, data, _ = sample
    modified = dict(data)
    modified["fact_inventory_snapshot"] = data["fact_inventory_snapshot"].copy()
    modified["fact_inventory_snapshot"].loc[96, "on_hand_units"] += 3
    assert len(reconcile_ledger(modified)) >= 1
    modified = dict(data)
    modified["fact_sales"] = data["fact_sales"].iloc[1:].copy()
    _, issues = normalize_and_validate(modified)
    assert "Missing sales observation" in set(issues["problem"])


def test_orphan_sku_and_missing_cost_fail_with_location(sample):
    _, data, _ = sample
    changed = dict(data)
    changed["dim_items"] = data["dim_items"].copy()
    changed["dim_items"].loc[0, "unit_cost_usd"] = None
    changed["fact_sales"] = data["fact_sales"].copy()
    changed["fact_sales"].loc[0, "sku_id"] = "NOT-A-SKU"
    _, issues = normalize_and_validate(changed)
    assert {"Orphan SKU", "Missing or nonpositive price"}.issubset(set(issues["problem"]))
    assert (issues["row"] == 2).any()


def test_duplicate_invalid_date_and_fractional_cent_fail_visibly(sample):
    _, data, _ = sample
    changed = dict(data)
    changed["fact_sales"] = pd.concat([data["fact_sales"], data["fact_sales"].iloc[[0]]])
    changed["fact_open_purchase_orders"] = data["fact_open_purchase_orders"].copy()
    changed["fact_open_purchase_orders"].loc[0, "order_date"] = "not-a-date"
    changed["dim_items"] = data["dim_items"].copy()
    changed["dim_items"].loc[0, "unit_cost_usd"] = 8.001
    _, issues = normalize_and_validate(changed)
    assert {"Duplicate key", "Invalid date", "Too many decimal places"}.issubset(set(issues["problem"]))
    assert issues["action"].notna().all()


def test_purchase_receipts_cancellations_and_overdue_asof():
    orders = pd.DataFrame(
        [
            ("P1", "SKU", "CDC", date(2026, 9, 1), date(2026, 9, 15), 100, 20, date(2026, 9, 17)),
            ("P2", "SKU", "CDC", date(2026, 9, 1), date(2026, 9, 25), 60, 60, date(2026, 9, 20)),
            ("P3", "SKU", "CDC", date(2026, 9, 19), date(2026, 10, 1), 40, 0, None),
        ],
        columns=[
            "po_id",
            "sku_id",
            "destination_loc",
            "order_date",
            "expected_delivery_date",
            "qty_ordered",
            "qty_cancelled",
            "cancellation_date",
        ],
    )
    movements = pd.DataFrame(
        [(date(2026, 9, 10), "SUPPLIER_RECEIPT", "P1", 30)],
        columns=["date", "movement_type", "reference_id", "delta_units"],
    )
    data = {"fact_open_purchase_orders": orders, "fact_movements": movements}
    assert _open_purchase_supply(data, date(2026, 9, 18)) == {
        "SKU": [(date(2026, 9, 19), 50), (date(2026, 9, 25), 60)]
    }
    assert _open_purchase_supply(data, date(2026, 9, 21)) == {
        "SKU": [(date(2026, 9, 22), 50), (date(2026, 10, 1), 40)]
    }


@pytest.mark.parametrize("budget", [0, 1_000_000, 5_000_000, 10_000_000])
def test_cash_cap_and_order_modifiers(sample, budget):
    _, data, forecast = sample
    result = decide(data, Scenario(budget_cents=budget), date(2026, 9, 18), forecast)
    assert result.spend_cents <= budget
    for order in result.buys.itertuples(index=False):
        assert valid_purchase(int(order.recommended_qty), int(order.moq_units), int(order.pack_multiple))
        assert order.cash_cents == order.recommended_qty * order.unit_cost_cents
    assert round_purchase(1, 50, 12) == 60
    assert not valid_purchase(48, 50, 12)


def test_transfer_source_conservation(sample):
    _, data, forecast = sample
    result = decide(data, Scenario(), date(2026, 9, 18), forecast)
    for sku, group in result.transfers.groupby("sku_id"):
        source = result.position[
            (result.position["sku_id"] == sku) & (result.position["location_id"] == "CDC")
        ].iloc[0]
        assert group["recommended_qty"].sum() <= source["on_hand_units"] - source["allocated_units"]


def test_future_sales_do_not_change_asof_decision(sample):
    _, data, forecast = sample
    changed = dict(data)
    changed["fact_sales"] = data["fact_sales"].copy()
    future = changed["fact_sales"]["date"] > date(2026, 9, 18)
    changed["fact_sales"].loc[future, "units_sold"] = 500
    original = decide(data, Scenario(), date(2026, 9, 18), forecast)
    altered = decide(changed, Scenario(), date(2026, 9, 18), forecast)
    assert original.buys[["sku_id", "recommended_qty"]].equals(altered.buys[["sku_id", "recommended_qty"]])


def test_export_matches_decision_and_opens(sample):
    _, data, forecast = sample
    result = decide(data, Scenario(), date(2026, 9, 18), forecast)
    workbook = openpyxl.load_workbook(BytesIO(excel_bytes(result)), read_only=True)
    assert set(workbook.sheetnames) == {
        "Summary",
        "Purchase Recommendations",
        "Transfer Recommendations",
        "Deferred Review Items",
        "Scenario Comparison",
        "Assumptions",
    }
    table = result_tables(result)["Purchase Recommendations"]
    assert round(table["total_cash_required_usd"].sum() * 100) == result.spend_cents
    assert workbook["Purchase Recommendations"].max_row == len(table) + 1


def test_replay_budget_and_demand_conservation(sample):
    generated, data, forecast = sample
    summary = replay(
        data, generated["truth_demand"], forecast, Scenario(budget_cents=2_500_000), date(2026, 3, 31)
    )
    assert summary["max_monthly_spend_usd"] <= 25_000
    assert summary["fulfilled_units"] + summary["lost_units"] == summary["demand_units"]
    assert summary["replay_days"] == 183
